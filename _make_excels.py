# -*- coding: utf-8 -*-
"""依網站實際結構生成「網站地圖.xlsx」與「網站變數與網站導覽表.xlsx」"""
import sys, os, unicodedata
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import core
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
from openpyxl.utils import get_column_letter

OUT = os.path.dirname(os.path.abspath(__file__))
BASE = core.SITE["url"].rstrip("/")
DATE = core.SITE.get("date", "2026-10-14")

DARK = "0E2A23"; GOLD = "C9A227"; IVORY = "F5F1E8"; LINE = "D8D2C4"
hdr_fill = PatternFill("solid", fgColor=DARK)
sub_fill = PatternFill("solid", fgColor=IVORY)
line_side = Side(style="thin", color=LINE)
border = Border(bottom=line_side)
hdr_font = Font(name="微軟雅黑", bold=True, color="FFFFFF", size=11)
body_font = Font(name="微軟雅黑", size=10)
gold_font = Font(name="微軟雅黑", size=10, color="0E2A23", bold=True)

def dw(v):
    return sum(2 if unicodedata.east_asian_width(ch) in ("F", "W") else 1 for ch in str(v or ""))

def fit(ws, min_col, max_col, min_row, max_row):
    for col in range(min_col, max_col + 1):
        L = get_column_letter(col)
        best = 0
        for row in range(min_row, max_row + 1):
            v = ws.cell(row, col).value
            if v is not None:
                best = max(best, dw(v))
        ws.column_dimensions[L].width = max(8, min(best + 3, 60))

def head_row(ws, row, headers, min_col=1):
    for i, h in enumerate(headers, min_col):
        c = ws.cell(row, i, h)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 22

# ---------- 頁面資料 ----------
pages = core.all_pages()
order_map = {".": 0, "brand": 1, "category": 2, "asset": 3, "criteria": 4,
             "service": 5, "feature": 6, "knowledge": 7, "appendix": 8}
def rank(p):
    od = p["out_dir"]
    if od.startswith("asset/"):
        return (3.5, od)
    return (order_map.get(od, 9), od)
items = sorted(pages.values(), key=lambda p: (rank(p), p["file"]))

# ================= Excel 1: 網站地圖 =================
wb = Workbook()
ws = wb.active; ws.title = "網站地圖"
headers = ["序號", "層級", "資料夾", "分類", "頁面標題", "H1", "檔名", "相對路徑", "網址", "更新日期"]
head_row(ws, 1, headers)
r = 2
for i, p in enumerate(items, 1):
    od = p["out_dir"]
    path = p["file"] if od in (".", "") else od + "/" + p["file"]
    level = 0 if od in (".", "") else od.count("/") + 1
    url = BASE + ("/index.html" if od in (".", "") and p["file"] == "index.html" else "/" + path)
    row = [i, level, od, p.get("section", ""), p.get("title", ""), p.get("h1", ""),
           p["file"], path, url, p.get("date", DATE)]
    for j, v in enumerate(row, 1):
        c = ws.cell(r, j, v); c.font = body_font
        c.alignment = Alignment(vertical="center")
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
    r += 1
fit(ws, 1, len(headers), 1, r - 1)
ws.freeze_panes = "A2"
# 統計區
sr = r + 1
ws.cell(sr, 1, "統計").font = gold_font
ws.cell(sr, 1).fill = sub_fill
stats = [
    ("頁面總數", len(items)),
    ("資料夾數（含根）", len({p["out_dir"] for p in items})),
    ("深度專區", "9 資產 × 22 篇 = 198 篇 + 9 總覽"),
    ("50 大功能", "50 頁"),
    ("知識體系", "11 頁"),
    ("附錄", "7 頁"),
]
for k, (label, val) in enumerate(stats):
    a = ws.cell(sr + k, 1, label); b = ws.cell(sr + k, 2, val)
    a.font = gold_font; b.font = body_font
    a.fill = sub_fill
wb.save(os.path.join(OUT, "網站地圖.xlsx"))

# ================= Excel 2: 網站變數與網站導覽表 =================
wb2 = Workbook()
ws1 = wb2.active; ws1.title = "網站變數"
head_row(ws1, 1, ["變數名", "值", "說明"])
vars_rows = [
    ("品牌名稱", "9TURN 九回錦", "硬資產典藏平台"),
    ("品牌英文名", "9TURN", "網域 9turn.com"),
    ("品牌定位", "只玩有價值的資產", "不玩短線炒作，只經營底層稀缺、可穿越週期的硬資產"),
    ("品牌口號", "九回見錦，價值歷循環而不減", "時間輪轉的價值哲學"),
    ("內容語言", "繁體中文（zh-Hant）", "台灣用語，UTF-8 編碼"),
    ("站點網址", BASE, "GitHub Pages 部署基址"),
    ("建置日期", DATE, "lastmod 基準"),
    ("技術棧", "純 HTML / CSS / JS（無後端）", "可直接部署 GitHub Pages"),
    ("字體", "Noto Serif TC（標題）＋ Noto Sans TC（內文）", "自托管鏡像載入"),
    ("主題色", "深黛青 #0E2A23／鎏金 #C9A227／象牙 #F5F1E8", "典藏質感"),
    ("圖片格式", "WebP", "lazy loading、語意化 alt"),
    ("頁面總數", len(items), "含 404 與 sitemap"),
    ("架構", "大主題＝子資料夾、每資料夾含 index.html", "搬移資料夾仍獨立成站"),
    ("導覽", "MegaMenu 巨型導覽（2 層內）", "＋ Breadcrumb／Sidebar＋Treemap／TOC"),
    ("搜尋", "全站搜尋 autocomplete＋Faceted 多面向篩選", "window.SITE_INDEX 內嵌索引"),
    ("收藏", "收藏清單（localStorage）", "右下角浮動導航＋pagination"),
    ("結構化資料", "JSON-LD：Organization／WebSite／BreadcrumbList／Article／FAQPage", "≥3 種"),
    ("SEO", "OG=9 完整設定、每頁 title／description／canonical／og:title", "＋ sitemap.xml／robots.txt／manifest／404"),
    ("檢核", "verify.py 三項批次掃描", "SEO 元素／斷鏈／本機檔案比對"),
    ("批次更新", "python core.py 全站重建", "改 content.py／deep_data 後一鍵重建"),
]
r = 2
for name, val, note in vars_rows:
    ws1.cell(r, 1, name).font = gold_font
    ws1.cell(r, 2, val).font = body_font
    ws1.cell(r, 3, note).font = body_font
    ws1.cell(r, 1).alignment = Alignment(vertical="center")
    ws1.cell(r, 2).alignment = Alignment(vertical="center")
    ws1.cell(r, 3).alignment = Alignment(vertical="center")
    r += 1
fit(ws1, 1, 3, 1, r - 1)
ws1.freeze_panes = "A2"

ws2 = wb2.create_sheet("網站導覽")
head_row(ws2, 1, ["頂層導覽", "子項", "目標路徑", "備註"])
r = 2
for title, url, children in core.NAV:
    ws2.cell(r, 1, title).font = gold_font
    ws2.cell(r, 2, "（總覽）").font = body_font
    ws2.cell(r, 3, url).font = body_font
    ws2.cell(r, 4, "MegaMenu 頂層入口").font = body_font
    ws2.cell(r, 1).alignment = Alignment(vertical="center")
    r += 1
    for child in children:
        if not child or len(child) < 2:
            continue
        if child[0] == "__head__":
            ws2.cell(r, 2, "— " + child[1] + " —").font = gold_font
            ws2.cell(r, 4, "分組標題").font = body_font
        else:
            ws2.cell(r, 2, child[0]).font = body_font
            ws2.cell(r, 3, child[1]).font = body_font
            ws2.cell(r, 4, "").font = body_font
        r += 1
fit(ws2, 1, 4, 1, r - 1)
ws2.freeze_panes = "A2"
wb2.save(os.path.join(OUT, "網站變數與網站導覽表.xlsx"))
print("OK: 網站地圖.xlsx / 網站變數與網站導覽表.xlsx")
print("頁面數:", len(items))
